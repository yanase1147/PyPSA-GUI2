"""component_templates シートの読み書きと PyPSA コンポーネントへの展開。

■ シートのレイアウト（人間が直接編集することを前提とした表形式）

    行1  タイトル + 凡例（結合セル）
    行2  グループ見出し（結合セル）      … 「接続先バス」「貯蔵パラメータ」など
    行3  日本語ラベル                   … 人間が読む列名
    行4  内部キー（機械可読）            … 読込はこの行を基準に列位置を決める
    行5〜 データ行                       … 1行 = 1構成要素 = PyPSA の Bus/Store/Link/Generator 1個

同じテンプレートに属する構成要素は連続した行に並べ、2行目以降の
「分類 / テンプレート名 / テンプレート説明」は空欄にできる（直前の行を引き継ぐ）。

■ セルの意味（ユーザー入力項目とシステム固定値の区別）

    数値・文字列 … システム側で固定する値（fixed_params）
    ``★入力``     … 設備ごとにユーザーが入力する項目（exposed_params）
    空欄          … 未使用（PyPSA の既定値のまま）

■ バス参照の書き方

    ``area:AC`` / ``area:hydrogen``  … そのエリアの共通バス（``エリア:`` も可）
    ``internal:store_1``             … 同じテンプレート内の Bus 要素（``内部:`` / ``#`` も可）
    ``AC`` のような裸のキャリア名     … ``area:AC`` として解釈

旧フォーマット（1セルに JSON を詰め込んだ ``json_data`` 形式）もそのまま読める。
"""
from __future__ import annotations

import json
from dataclasses import dataclass
from typing import Any, Callable, Dict, List, Optional, Sequence, Tuple

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .models import (
    AVAILABLE_EXPOSED_PARAMS,
    ComponentTemplate,
    CustomComponentInstance,
    MULTI_CARRIERS,
    SubComponentDef,
)

SHEET_NAME = "component_templates"

#: ユーザー入力項目を表すマーカー（書き出し時はこの文字列を使う）
EXPOSED_MARKER = "★入力"

#: 読込時にユーザー入力項目とみなす表記
_EXPOSED_TOKENS = {
    "★", "★入力", "☆", "入力", "要入力", "ユーザー入力", "ユーザ入力",
    "user", "input", "user_input", "expose", "exposed",
}

_TRUE_TOKENS = {"true", "1", "yes", "y", "はい", "○", "◯", "有効", "on"}
_FALSE_TOKENS = {"false", "0", "no", "n", "いいえ", "×", "✕", "無効", "off", ""}

#: PyPSA 種別ごとに使うバス参照スロット
BUS_SLOTS: Dict[str, Tuple[str, ...]] = {
    "Bus": (),
    "Store": ("bus",),
    "Generator": ("bus",),
    "Link": ("bus0", "bus1", "bus2"),
}

COMPONENT_TYPES = ("Bus", "Store", "Link", "Generator")

#: 既知のキャリア（裸で書かれたバス参照を ``area:`` に補完するために使う）
_BARE_CARRIERS = {"AC", "DC", "other", *MULTI_CARRIERS}


# ══════════════════════════════════════════════════════════════════════
# 列定義
# ══════════════════════════════════════════════════════════════════════

@dataclass(frozen=True)
class _Col:
    key: str        # 内部キー（行4に書かれる）
    label: str      # 日本語ラベル（行3に書かれる）
    group: str      # グループ見出し（行2に書かれる）
    kind: str       # "meta" | "flag" | "bus" | "param" | "extra" | "layout"
    width: int = 14


_G_TMPL = "① テンプレート（分類・名称）"
_G_SUB = "② 構成要素（PyPSA 部品）"
_G_BUS = "③ 接続先バス"
_G_COMMON = "④ 共通"
_G_POWER = "⑤ 出力側パラメータ（Link / Generator）"
_G_STORE = "⑥ 貯蔵パラメータ（Store）"
_G_COST = "⑦ コスト・年次"
_G_ADV = "⑧ 上級者向け"
_G_AUTO = "⑨ 自動生成（編集不要）"

COLUMNS: Tuple[_Col, ...] = (
    _Col("category",             "分類",                        _G_TMPL, "meta", 14),
    _Col("template_name",        "テンプレート名",               _G_TMPL, "meta", 18),
    _Col("template_description", "テンプレート説明",             _G_TMPL, "meta", 26),
    _Col("enabled",              "有効",                        _G_TMPL, "flag", 8),

    _Col("sub_id",               "要素ID",                      _G_SUB, "meta", 12),
    _Col("component_type",       "PyPSA種別",                   _G_SUB, "meta", 12),
    _Col("name_template",        "要素名（{name}=設備名）",       _G_SUB, "meta", 24),
    _Col("sub_description",      "要素の役割",                   _G_SUB, "meta", 26),

    _Col("bus",                  "接続バス（Store/Gen）",         _G_BUS, "bus", 18),
    _Col("bus0",                 "入力バス bus0（Link）",         _G_BUS, "bus", 18),
    _Col("bus1",                 "出力バス bus1（Link）",         _G_BUS, "bus", 18),
    _Col("bus2",                 "副出力バス bus2（Link）",       _G_BUS, "bus", 18),

    _Col("carrier",              "キャリア",                     _G_COMMON, "param", 12),

    _Col("efficiency",           "効率",                        _G_POWER, "param", 10),
    _Col("efficiency2",          "効率2（副出力）",              _G_POWER, "param", 12),
    _Col("p_nom",                "出力容量 [MW]",                _G_POWER, "param", 13),
    _Col("p_nom_extendable",     "出力容量を最適化",              _G_POWER, "param", 15),
    _Col("p_nom_max",            "最大出力容量 [MW]",            _G_POWER, "param", 15),
    _Col("p_min_pu",             "最小出力率",                   _G_POWER, "param", 11),
    _Col("p_max_pu",             "最大出力率",                   _G_POWER, "param", 11),

    _Col("e_nom",                "蓄電容量 [MWh]",               _G_STORE, "param", 14),
    _Col("e_nom_extendable",     "蓄電容量を最適化",              _G_STORE, "param", 15),
    _Col("e_nom_max",            "最大蓄電容量 [MWh]",           _G_STORE, "param", 16),
    _Col("e_min_pu",             "最小充電率",                   _G_STORE, "param", 11),
    _Col("e_max_pu",             "最大充電率",                   _G_STORE, "param", 11),
    _Col("e_initial",            "初期蓄電量 [MWh]",             _G_STORE, "param", 14),
    _Col("standing_loss",        "自己放電率 [/h]",              _G_STORE, "param", 13),
    _Col("cyclic_state_of_charge", "周期境界条件",                _G_STORE, "param", 13),

    _Col("capital_cost",         "建設単価 [通貨/MW・MWh]",       _G_COST, "param", 18),
    _Col("marginal_cost",        "変動費 [通貨/MWh]",            _G_COST, "param", 15),
    _Col("build_year",           "建設年",                      _G_COST, "param", 10),

    _Col("extra_params_json",    "その他パラメータ(JSON)",        _G_ADV, "extra", 24),

    _Col("pos_x",                "図X座標",                     _G_AUTO, "layout", 10),
    _Col("pos_y",                "図Y座標",                     _G_AUTO, "layout", 10),
)

_COL_BY_KEY: Dict[str, _Col] = {c.key: c for c in COLUMNS}

#: 表の列として用意されているパラメータ
PARAM_COLUMNS: Tuple[str, ...] = tuple(c.key for c in COLUMNS if c.kind == "param")

#: パラメータの型（表記ゆれを吸収して PyPSA に渡せる型へ変換する）
_PARAM_TYPES: Dict[str, type] = {
    "carrier": str,
    "efficiency": float, "efficiency2": float,
    "p_nom": float, "p_nom_max": float, "p_min_pu": float, "p_max_pu": float,
    "p_nom_extendable": bool,
    "e_nom": float, "e_nom_max": float, "e_min_pu": float, "e_max_pu": float,
    "e_initial": float, "standing_loss": float,
    "e_nom_extendable": bool, "cyclic_state_of_charge": bool,
    "capital_cost": float, "marginal_cost": float,
    "build_year": int,
}

_TITLE = "component_templates … 複合コンポーネント（蓄電池・水素貯蔵・揚水発電など）の定義表"
_LEGEND = (
    f"【凡例】 数値/文字を入れた欄＝システム固定値、"
    f"「{EXPOSED_MARKER}」＝設備ごとにユーザーが入力する項目、空欄＝未使用（PyPSA既定値）。  "
    "【バス参照】 area:AC / area:hydrogen＝エリア共通バス、internal:要素ID＝テンプレート内部バス。  "
    "【行の続き】 分類・テンプレート名・テンプレート説明が空欄の行は、直前の行と同じテンプレートの構成要素として扱われます。"
)

# ── 見た目 ────────────────────────────────────────────────────────────
_GROUP_COLORS: Dict[str, Tuple[str, str]] = {
    # group: (見出しの塗り, データ列の薄い塗り)
    _G_TMPL:   ("1F4E79", "D9E2F3"),
    _G_SUB:    ("2E75B6", "DEEBF7"),
    _G_BUS:    ("548235", "E2EFDA"),
    _G_COMMON: ("595959", "EDEDED"),
    _G_POWER:  ("BF8F00", "FFF2CC"),
    _G_STORE:  ("9E480E", "FBE5D6"),
    _G_COST:   ("C55A11", "F8CBAD"),
    _G_ADV:    ("7030A0", "E4DFEC"),
    _G_AUTO:   ("808080", "F2F2F2"),
}

_THIN = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

TITLE_ROW = 1
GROUP_ROW = 2
LABEL_ROW = 3
KEY_ROW = 4
FIRST_DATA_ROW = 5


# ══════════════════════════════════════════════════════════════════════
# セル値のパース
# ══════════════════════════════════════════════════════════════════════

def _s(v: Any) -> str:
    return "" if v is None else str(v).strip()


def _is_exposed(v: Any) -> bool:
    """セルが「ユーザーが設備ごとに入力する項目」マーカーかどうか。"""
    t = _s(v)
    if not t:
        return False
    return t.lower() in _EXPOSED_TOKENS or t.startswith("★") or t.startswith("☆")


def _to_bool(v: Any, default: bool = False) -> bool:
    if isinstance(v, bool):
        return v
    t = _s(v).lower()
    if t in _TRUE_TOKENS:
        return True
    if t in _FALSE_TOKENS:
        return False
    return default


def _coerce(key: str, v: Any) -> Any:
    """セル値を PyPSA に渡せる型へ変換する。変換できなければ元の文字列。"""
    typ = _PARAM_TYPES.get(key, str)
    if typ is bool:
        return _to_bool(v)
    if typ is str:
        return _s(v)
    try:
        if typ is int:
            return int(float(_s(v).replace(",", "")))
        return float(_s(v).replace(",", ""))
    except (TypeError, ValueError):
        return _s(v)


def normalize_bus_ref(raw: Any) -> str:
    """バス参照セルを正規形（``area:X`` / ``internal:Y``）へ変換する。"""
    t = _s(raw)
    if not t:
        return ""
    for prefix in ("エリア:", "エリア：", "area："):
        if t.startswith(prefix):
            t = "area:" + t[len(prefix):].strip()
            break
    for prefix in ("内部:", "内部：", "internal：", "sub:", "要素:"):
        if t.startswith(prefix):
            t = "internal:" + t[len(prefix):].strip()
            break
    if t.startswith("#"):
        t = "internal:" + t[1:].strip()
    if t.startswith("@"):
        t = "area:" + t[1:].strip()
    if t.startswith("area:") or t.startswith("internal:"):
        head, _, tail = t.partition(":")
        return f"{head}:{tail.strip()}"
    if t in _BARE_CARRIERS:
        return f"area:{t}"
    return t


# ══════════════════════════════════════════════════════════════════════
# 読込
# ══════════════════════════════════════════════════════════════════════

def load_component_templates(source) -> List[ComponentTemplate]:
    """``component_templates`` シートを読み込み ComponentTemplate のリストを返す。

    Parameters
    ----------
    source :
        Excel ファイルのパス / ``openpyxl`` の Workbook / Worksheet。

    新フォーマット（表形式）と旧フォーマット（``json_data`` 1セル）の
    どちらでも読み込める。シートが無い場合は空リストを返す。
    """
    ws = _resolve_worksheet(source)
    if ws is None:
        return []
    if _looks_like_legacy_json(ws):
        return _load_legacy_json(ws)
    return _load_table(ws)


def _resolve_worksheet(source):
    if source is None:
        return None
    if hasattr(source, "iter_rows") and not hasattr(source, "sheetnames"):
        return source                                   # Worksheet
    wb = source
    if isinstance(source, str):
        wb = openpyxl.load_workbook(source, data_only=True)
    if hasattr(wb, "sheetnames"):
        return wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else None
    return None


def _looks_like_legacy_json(ws) -> bool:
    return _s(ws.cell(row=1, column=1).value) == "json_data"


def _load_legacy_json(ws) -> List[ComponentTemplate]:
    """旧フォーマット（1セルに JSON）を読む。"""
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    if not rows or not rows[0] or not rows[0][0]:
        return []
    try:
        payload = json.loads(rows[0][0])
    except (ValueError, TypeError):
        return []
    templates: List[ComponentTemplate] = []
    for td in payload:
        subs = []
        for sd in td.get("sub_components", []):
            known = {f: sd[f] for f in SubComponentDef.__dataclass_fields__ if f in sd}
            subs.append(SubComponentDef(**known))
        templates.append(ComponentTemplate(
            name=td.get("name", ""),
            description=td.get("description", ""),
            sub_components=subs,
            category=td.get("category", ""),
        ))
    return templates


def _find_key_row(ws, max_scan: int = 20) -> Optional[Tuple[int, Dict[str, int]]]:
    """内部キー行を探し、(行番号, {キー: 列番号}) を返す。"""
    for row in ws.iter_rows(min_row=1, max_row=max_scan):
        keys = {}
        for cell in row:
            k = _s(cell.value)
            if k in _COL_BY_KEY and k not in keys:
                keys[k] = cell.column
        if "sub_id" in keys and "component_type" in keys and "template_name" in keys:
            return row[0].row, keys
    return None


def _load_table(ws) -> List[ComponentTemplate]:
    found = _find_key_row(ws)
    if not found:
        return []
    key_row, col_of = found

    templates: List[ComponentTemplate] = []
    by_name: Dict[str, ComponentTemplate] = {}
    carry = {"category": "", "template_name": "", "template_description": ""}

    for row in ws.iter_rows(min_row=key_row + 1):
        cells = {k: row[c - 1].value for k, c in col_of.items() if c - 1 < len(row)}
        if all(_s(v) == "" for v in cells.values()):
            continue

        # テンプレート名が書かれている行は新しいテンプレートの開始。
        # 空欄の行は直前のテンプレートの構成要素として引き継ぐ。
        raw_name = _s(cells.get("template_name"))
        if raw_name:
            carry["template_name"] = raw_name
            carry["category"] = _s(cells.get("category"))
            carry["template_description"] = _s(cells.get("template_description"))
        else:
            for k in ("category", "template_description"):
                v = _s(cells.get(k))
                if v:
                    carry[k] = v
        tmpl_name = carry["template_name"]
        if not tmpl_name:
            continue

        tmpl = by_name.get(tmpl_name)
        if tmpl is None:
            tmpl = ComponentTemplate(
                name=tmpl_name,
                description=carry["template_description"],
                category=carry["category"],
            )
            by_name[tmpl_name] = tmpl
            templates.append(tmpl)
        else:
            # 後続行で説明・分類が書かれていれば補完する
            tmpl.description = tmpl.description or carry["template_description"]
            tmpl.category = tmpl.category or carry["category"]

        if "enabled" in cells and not _to_bool(cells.get("enabled"), True):
            continue

        sub = _row_to_sub(cells, tmpl)
        if sub is not None:
            tmpl.sub_components.append(sub)

    return templates


def _row_to_sub(cells: Dict[str, Any], tmpl: ComponentTemplate) -> Optional[SubComponentDef]:
    ctype = _s(cells.get("component_type"))
    if not ctype:
        return None
    # 「store」「STORE」などの表記ゆれを吸収
    ctype = next((c for c in COMPONENT_TYPES if c.lower() == ctype.lower()), ctype)

    sub_id = _s(cells.get("sub_id")) or f"{ctype.lower()}_{len(tmpl.sub_components) + 1}"
    name_template = _s(cells.get("name_template")) or "{name}"

    fixed: Dict[str, Any] = {}
    exposed: List[str] = []
    allowed = AVAILABLE_EXPOSED_PARAMS.get(ctype, [])
    for key in PARAM_COLUMNS:
        raw = cells.get(key)
        if _s(raw) == "":
            continue
        if _is_exposed(raw):
            if key in allowed:
                exposed.append(key)
            continue
        fixed[key] = _coerce(key, raw)

    extra = _s(cells.get("extra_params_json"))
    if extra:
        try:
            parsed = json.loads(extra)
        except ValueError:
            parsed = None
        if isinstance(parsed, dict):
            for name in parsed.pop("__exposed__", []) or []:
                if name not in exposed:
                    exposed.append(name)
            fixed.update(parsed)

    bus_connections: Dict[str, str] = {}
    for slot in BUS_SLOTS.get(ctype, ()):
        ref = normalize_bus_ref(cells.get(slot))
        if ref:
            bus_connections[slot] = ref

    return SubComponentDef(
        sub_id=sub_id,
        component_type=ctype,
        name_template=name_template,
        fixed_params=fixed,
        exposed_params=exposed,
        bus_connections=bus_connections,
        pos_x=_num(cells.get("pos_x")),
        pos_y=_num(cells.get("pos_y")),
        description=_s(cells.get("sub_description")),
    )


def _num(v: Any, default: float = 0.0) -> float:
    try:
        return float(_s(v))
    except (TypeError, ValueError):
        return default


# ══════════════════════════════════════════════════════════════════════
# 書き出し
# ══════════════════════════════════════════════════════════════════════

def template_rows(templates: Sequence[ComponentTemplate]) -> List[Dict[str, Any]]:
    """テンプレート群をシートの行（キー→値の辞書）へ展開する。"""
    rows: List[Dict[str, Any]] = []
    for tmpl in templates:
        subs = tmpl.sub_components or [None]
        for i, sub in enumerate(subs):
            row: Dict[str, Any] = {}
            if i == 0:
                row["category"] = tmpl.category
                row["template_name"] = tmpl.name
                row["template_description"] = tmpl.description
            row["enabled"] = True
            if sub is None:
                rows.append(row)
                continue

            row["sub_id"] = sub.sub_id
            row["component_type"] = sub.component_type
            row["name_template"] = sub.name_template
            row["sub_description"] = sub.description
            for slot in BUS_SLOTS.get(sub.component_type, ()):
                ref = sub.bus_connections.get(slot, "")
                if ref:
                    row[slot] = ref
            # bus_connections に想定外のスロットがあっても落とさない
            for slot, ref in sub.bus_connections.items():
                if slot in _COL_BY_KEY and slot not in row and ref:
                    row[slot] = ref

            leftovers = dict(sub.fixed_params)
            for key in PARAM_COLUMNS:
                if key in sub.exposed_params:
                    row[key] = EXPOSED_MARKER
                    leftovers.pop(key, None)
                elif key in leftovers:
                    row[key] = leftovers.pop(key)
            # 列を持たない exposed パラメータも失わないよう JSON 側に退避
            extra_exposed = [p for p in sub.exposed_params if p not in PARAM_COLUMNS]
            if leftovers or extra_exposed:
                blob: Dict[str, Any] = dict(leftovers)
                if extra_exposed:
                    blob["__exposed__"] = extra_exposed
                row["extra_params_json"] = json.dumps(blob, ensure_ascii=False)

            # Excel の有効桁で丸め差が出ないよう固定小数にする（表示用座標）
            row["pos_x"] = round(float(sub.pos_x or 0.0), 2)
            row["pos_y"] = round(float(sub.pos_y or 0.0), 2)
            rows.append(row)
    return rows


def write_component_templates_sheet(wb, templates: Sequence[ComponentTemplate]):
    """ワークブックに新フォーマットの ``component_templates`` シートを作成する。"""
    if SHEET_NAME in wb.sheetnames:
        del wb[SHEET_NAME]
    ws = wb.create_sheet(SHEET_NAME)

    n_cols = len(COLUMNS)
    last_letter = get_column_letter(n_cols)

    # ── 行1: タイトル + 凡例 ───────────────────────────────────────
    ws.merge_cells(f"A{TITLE_ROW}:{last_letter}{TITLE_ROW}")
    title = ws.cell(row=TITLE_ROW, column=1, value=f"{_TITLE}\n{_LEGEND}")
    title.font = Font(bold=True, size=10, color="1F4E79")
    title.alignment = Alignment(vertical="center", wrap_text=True)
    title.fill = PatternFill("solid", start_color="F2F7FB")
    ws.row_dimensions[TITLE_ROW].height = 48

    # ── 行2: グループ見出し（連続する同一グループを結合） ─────────
    start = 1
    for i in range(1, n_cols + 2):
        cur = COLUMNS[i - 1].group if i <= n_cols else None
        prev = COLUMNS[start - 1].group
        if cur == prev:
            continue
        head, _ = _GROUP_COLORS.get(prev, ("4472C4", "FFFFFF"))
        if i - 1 > start:
            ws.merge_cells(start_row=GROUP_ROW, start_column=start,
                           end_row=GROUP_ROW, end_column=i - 1)
        for c in range(start, i):
            cell = ws.cell(row=GROUP_ROW, column=c)
            cell.fill = PatternFill("solid", start_color=head)
            cell.border = _BORDER
        cell = ws.cell(row=GROUP_ROW, column=start, value=prev)
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        start = i
    ws.row_dimensions[GROUP_ROW].height = 20

    # ── 行3: 日本語ラベル / 行4: 内部キー ──────────────────────────
    for idx, col in enumerate(COLUMNS, start=1):
        head, body = _GROUP_COLORS.get(col.group, ("4472C4", "FFFFFF"))

        label = ws.cell(row=LABEL_ROW, column=idx, value=col.label)
        label.font = Font(bold=True, size=9)
        label.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        label.fill = PatternFill("solid", start_color=body)
        label.border = _BORDER

        key = ws.cell(row=KEY_ROW, column=idx, value=col.key)
        key.font = Font(size=8, color="808080", italic=True)
        key.alignment = Alignment(horizontal="center")
        key.fill = PatternFill("solid", start_color="FFFFFF")
        key.border = _BORDER

        ws.column_dimensions[get_column_letter(idx)].width = col.width
    ws.row_dimensions[LABEL_ROW].height = 32

    # ── 行5〜: データ ──────────────────────────────────────────────
    rows = template_rows(templates)
    for r_off, row in enumerate(rows):
        r = FIRST_DATA_ROW + r_off
        for idx, col in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=r, column=idx, value=row.get(col.key))
            cell.border = _BORDER
            cell.font = Font(size=9)
            if col.kind == "layout":
                cell.font = Font(size=8, color="A6A6A6")
            elif _s(cell.value) == EXPOSED_MARKER:
                cell.font = Font(size=9, bold=True, color="C00000")
                cell.fill = PatternFill("solid", start_color="FFF2F2")
                cell.alignment = Alignment(horizontal="center")
        if _s(row.get("template_name")):
            for idx in range(1, n_cols + 1):
                ws.cell(row=r, column=idx).border = Border(
                    left=_THIN, right=_THIN, bottom=_THIN,
                    top=Side(style="medium", color="808080"))

    _add_validations(ws, len(rows))
    ws.freeze_panes = ws.cell(row=FIRST_DATA_ROW, column=_col_index("sub_id"))
    ws.sheet_view.showGridLines = False
    return ws


def _col_index(key: str) -> int:
    for i, col in enumerate(COLUMNS, start=1):
        if col.key == key:
            return i
    return 1


def _add_validations(ws, n_rows: int) -> None:
    """入力ミスを減らすためのドロップダウンを設定する。"""
    last = max(FIRST_DATA_ROW + n_rows + 50, FIRST_DATA_ROW + 50)

    def add(key: str, options: Sequence[str], prompt: str) -> None:
        idx = _col_index(key)
        letter = get_column_letter(idx)
        dv = DataValidation(
            type="list",
            formula1='"' + ",".join(options) + '"',
            allow_blank=True,
            showErrorMessage=False,
        )
        dv.promptTitle = key
        dv.prompt = prompt
        dv.showInputMessage = True
        ws.add_data_validation(dv)
        dv.add(f"{letter}{FIRST_DATA_ROW}:{letter}{last}")

    add("component_type", COMPONENT_TYPES, "この行が生成する PyPSA の部品種別")
    add("enabled", ("TRUE", "FALSE"), "FALSE にするとこの行は読み飛ばされます")
    add("carrier", ("AC", "DC", *MULTI_CARRIERS, "other"), "エネルギーキャリア")
    for k in ("p_nom_extendable", "e_nom_extendable", "cyclic_state_of_charge"):
        add(k, ("TRUE", "FALSE", EXPOSED_MARKER),
            f"TRUE/FALSE で固定、{EXPOSED_MARKER} で設備ごとの入力項目")


# ══════════════════════════════════════════════════════════════════════
# PyPSA への展開
# ══════════════════════════════════════════════════════════════════════

def resolve_bus_ref(ref: str, area: str, instance_name: str,
                    tmpl: ComponentTemplate,
                    sub_name_map: Optional[Dict[str, str]] = None,
                    bus_name_fn: Optional[Callable[[str, str], str]] = None) -> str:
    """``area:`` / ``internal:`` 参照を実際の PyPSA バス名へ変換する。"""
    ref = normalize_bus_ref(ref)
    if ref.startswith("area:"):
        carrier = ref[5:]
        if bus_name_fn is not None:
            return bus_name_fn(area, carrier)
        return area if carrier in ("AC", "", None) else f"{area}-{carrier}"
    if ref.startswith("internal:"):
        sub_id = ref[9:]
        if sub_name_map and sub_id in sub_name_map:
            return sub_name_map[sub_id]
        sub = next((s for s in tmpl.sub_components if s.sub_id == sub_id), None)
        if sub:
            return sub.name_template.replace("{name}", instance_name)
    return ref


def build_sub_name_map(tmpl: ComponentTemplate, instance_name: str,
                       name_fn: Optional[Callable[[SubComponentDef, str], str]] = None,
                       ) -> Dict[str, str]:
    """要素ID → 実際のコンポーネント名 の対応表を作る。"""
    out: Dict[str, str] = {}
    for sub in tmpl.sub_components:
        base = sub.name_template.replace("{name}", instance_name)
        out[sub.sub_id] = name_fn(sub, base) if name_fn else base
    return out


def expand_instance(
    tmpl: ComponentTemplate,
    instance: CustomComponentInstance,
    *,
    sub_name_map: Optional[Dict[str, str]] = None,
    bus_name_fn: Optional[Callable[[str, str], str]] = None,
    param_hook: Optional[Callable[[SubComponentDef, Dict[str, Any]], Dict[str, Any]]] = None,
) -> List[Tuple[str, str, Dict[str, Any]]]:
    """1つの設備インスタンスを PyPSA コンポーネントの列へ展開する。

    Returns
    -------
    ``[(component_type, name, params), ...]``
        ``params`` にはバス名まで解決済みのキーワード引数が入る。
        呼び出し側は ``n.add(component_type, name, **params)`` すればよい。

    Parameters
    ----------
    sub_name_map :
        要素ID → 実名。省略時は ``name_template`` から生成する。
    bus_name_fn :
        ``(エリア名, キャリア) -> バス名``。省略時は ``AC`` はエリア名、
        それ以外は ``{エリア}-{キャリア}``。
    param_hook :
        コンポーネント追加直前に params を加工する関数（年経費化など）。
    """
    if sub_name_map is None:
        sub_name_map = build_sub_name_map(tmpl, instance.name)

    out: List[Tuple[str, str, Dict[str, Any]]] = []
    for sub in tmpl.sub_components:
        name = sub_name_map.get(
            sub.sub_id, sub.name_template.replace("{name}", instance.name))

        params: Dict[str, Any] = dict(sub.fixed_params)
        for p_name in sub.exposed_params:
            key = f"{sub.sub_id}.{p_name}"
            if key in instance.param_values:
                params[p_name] = instance.param_values[key]

        for slot in BUS_SLOTS.get(sub.component_type, ()):
            ref = sub.bus_connections.get(slot)
            if ref:
                params[slot] = resolve_bus_ref(
                    ref, instance.area, instance.name, tmpl,
                    sub_name_map, bus_name_fn)

        if param_hook is not None:
            params = param_hook(sub, params)

        out.append((sub.component_type, name, params))
    return out


# ══════════════════════════════════════════════════════════════════════
# サンプル（新フォーマットの書き方の見本）
# ══════════════════════════════════════════════════════════════════════

def sample_templates() -> List[ComponentTemplate]:
    """記入例として使える代表的な複合コンポーネント定義を返す。"""
    return [
        ComponentTemplate(
            name="リチウムイオン蓄電池",
            category="蓄電池",
            description="充電Link＋Store＋放電LinkをAC母線に接続した標準的な蓄電池",
            sub_components=[
                SubComponentDef(
                    sub_id="bat_bus", component_type="Bus",
                    name_template="{name}-母線",
                    fixed_params={"carrier": "other"},
                    description="蓄電池の内部母線（システム自動生成）",
                    pos_x=0, pos_y=0),
                SubComponentDef(
                    sub_id="charge", component_type="Link",
                    name_template="{name}-充電",
                    fixed_params={"efficiency": 0.95},
                    exposed_params=["p_nom", "p_nom_extendable", "capital_cost"],
                    bus_connections={"bus0": "area:AC", "bus1": "internal:bat_bus"},
                    description="系統→蓄電池（充電効率95%）",
                    pos_x=-220, pos_y=0),
                SubComponentDef(
                    sub_id="cell", component_type="Store",
                    name_template="{name}-セル",
                    fixed_params={"carrier": "other", "standing_loss": 0.00002,
                                  "cyclic_state_of_charge": True, "e_max_pu": 1.0,
                                  "e_min_pu": 0.0},
                    exposed_params=["e_nom", "e_nom_extendable", "capital_cost"],
                    bus_connections={"bus": "internal:bat_bus"},
                    description="蓄電容量本体（自己放電0.002%/h）",
                    pos_x=0, pos_y=160),
                SubComponentDef(
                    sub_id="discharge", component_type="Link",
                    name_template="{name}-放電",
                    fixed_params={"efficiency": 0.95, "marginal_cost": 0.5},
                    exposed_params=["p_nom", "p_nom_extendable"],
                    bus_connections={"bus0": "internal:bat_bus", "bus1": "area:AC"},
                    description="蓄電池→系統（放電効率95%）",
                    pos_x=220, pos_y=0),
            ],
        ),
        ComponentTemplate(
            name="水素貯蔵施設",
            category="水素貯蔵",
            description="水素母線に直結する貯蔵タンク（電解槽・燃料電池は別テンプレート）",
            sub_components=[
                SubComponentDef(
                    sub_id="h2_tank", component_type="Store",
                    name_template="{name}-タンク",
                    fixed_params={"carrier": "hydrogen", "cyclic_state_of_charge": True,
                                  "standing_loss": 0.0},
                    exposed_params=["e_nom", "e_nom_extendable", "e_nom_max",
                                    "capital_cost", "marginal_cost"],
                    bus_connections={"bus": "area:hydrogen"},
                    description="水素貯蔵タンク本体",
                    pos_x=-40, pos_y=0),
            ],
        ),
        ComponentTemplate(
            name="水素製造施設（電解槽）",
            category="水素貯蔵",
            description="AC母線の電力で水素を製造し水素母線へ供給する電解槽",
            sub_components=[
                SubComponentDef(
                    sub_id="electrolyser", component_type="Link",
                    name_template="{name}-電解槽",
                    fixed_params={"efficiency": 0.7},
                    exposed_params=["p_nom", "p_nom_extendable", "p_nom_max",
                                    "capital_cost", "marginal_cost", "build_year"],
                    bus_connections={"bus0": "area:AC", "bus1": "area:hydrogen"},
                    description="電力→水素（効率70%）",
                    pos_x=-25, pos_y=50),
            ],
        ),
        ComponentTemplate(
            name="揚水発電所",
            category="揚水発電",
            description="上池を Store、揚水/発電を Link で表現した揚水発電所",
            sub_components=[
                SubComponentDef(
                    sub_id="upper_bus", component_type="Bus",
                    name_template="{name}-上池母線",
                    fixed_params={"carrier": "other"},
                    description="上池（水）を表す内部母線（システム自動生成）",
                    pos_x=0, pos_y=0),
                SubComponentDef(
                    sub_id="pump", component_type="Link",
                    name_template="{name}-揚水",
                    fixed_params={"efficiency": 0.87},
                    exposed_params=["p_nom", "p_nom_extendable", "capital_cost"],
                    bus_connections={"bus0": "area:AC", "bus1": "internal:upper_bus"},
                    description="系統→上池（ポンプ効率87%）",
                    pos_x=-220, pos_y=0),
                SubComponentDef(
                    sub_id="reservoir", component_type="Store",
                    name_template="{name}-上池",
                    fixed_params={"carrier": "other", "cyclic_state_of_charge": True,
                                  "standing_loss": 0.0},
                    exposed_params=["e_nom", "e_nom_extendable", "e_initial"],
                    bus_connections={"bus": "internal:upper_bus"},
                    description="上池の貯水量 [MWh 相当]",
                    pos_x=0, pos_y=160),
                SubComponentDef(
                    sub_id="turbine", component_type="Link",
                    name_template="{name}-発電",
                    fixed_params={"efficiency": 0.9, "marginal_cost": 0.0},
                    exposed_params=["p_nom", "p_nom_extendable", "capital_cost"],
                    bus_connections={"bus0": "internal:upper_bus", "bus1": "area:AC"},
                    description="上池→系統（水車効率90%）",
                    pos_x=220, pos_y=0),
            ],
        ),
    ]


def write_sample_workbook(path: str) -> str:
    """記入例だけを収めたワークブックを書き出す（レイアウト確認用）。"""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    write_component_templates_sheet(wb, sample_templates())
    wb.save(path)
    return path


# ══════════════════════════════════════════════════════════════════════
# 既存ブックの移行
# ══════════════════════════════════════════════════════════════════════

def migrate_workbook(path: str, out_path: Optional[str] = None) -> int:
    """既存の Excel の ``component_templates`` シートを新フォーマットへ書き換える。

    Returns
    -------
    int : 書き出したテンプレート数
    """
    templates = load_component_templates(path)
    wb = openpyxl.load_workbook(path)
    order = wb.sheetnames.index(SHEET_NAME) if SHEET_NAME in wb.sheetnames else None
    ws = write_component_templates_sheet(wb, templates)
    if order is not None:
        wb.move_sheet(ws, offset=order - wb.sheetnames.index(SHEET_NAME))
    wb.save(out_path or path)
    return len(templates)


if __name__ == "__main__":  # pragma: no cover
    import sys

    if len(sys.argv) < 2:
        print("usage:")
        print("  python -m src.component_templates <project.xlsx> [out.xlsx]"
              "   # 旧フォーマット → 新フォーマットへ移行")
        print("  python -m src.component_templates --sample <out.xlsx>"
              "        # 記入例ブックを書き出す")
        raise SystemExit(1)
    if sys.argv[1] == "--sample":
        dest = sys.argv[2] if len(sys.argv) > 2 else "component_templates_sample.xlsx"
        print("wrote sample:", write_sample_workbook(dest))
    else:
        dest = sys.argv[2] if len(sys.argv) > 2 else sys.argv[1]
        n = migrate_workbook(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None)
        print(f"migrated {n} template(s) in {dest}")
