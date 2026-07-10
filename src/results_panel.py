"""Results panel: capacity/generation/cost/CO₂ charts + multi-year comparison + Excel export."""
from __future__ import annotations

import datetime
import os
import re
import shutil
import tempfile
import traceback
from typing import Optional

import numpy as np
import matplotlib
matplotlib.use("QtAgg")
from matplotlib.backends.backend_qtagg import FigureCanvasQTAgg as FigureCanvas
from matplotlib.figure import Figure
from matplotlib.ticker import FuncFormatter
import matplotlib.pyplot as plt

from PyQt6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QTabWidget, QComboBox,
    QLabel, QPushButton, QFileDialog, QMessageBox, QScrollArea,
    QGroupBox, QTableWidget, QTableWidgetItem, QHeaderView,
    QSizePolicy, QSpinBox, QSplitter,
)
from PyQt6.QtCore import Qt

from .models import OptimizationResults, YearResult, CARRIER_COLORS, CF_CARRIERS
from .pypsa_runner import extract_year_results, extract_year_timeseries, _extract_multi_period_year


class _SummaryWindow(QWidget):
    """電源別サマリーを表示する独立ウィンドウ。"""

    def __init__(self, parent=None):
        super().__init__(parent, Qt.WindowType.Window)
        self.setWindowTitle(self.tr("電源別サマリー"))
        self.resize(680, 360)
        layout = QVBoxLayout(self)
        layout.setContentsMargins(6, 6, 6, 6)

        self.title_label = QLabel("")
        layout.addWidget(self.title_label)

        self.table = QTableWidget(0, 5)
        self.table.setHorizontalHeaderLabels([
            self.tr("電源種別"), self.tr("設備容量(MW)"),
            self.tr("発電量(MWh)"), self.tr("建設費(Currency/年)"), self.tr("運転費(Currency/年)")])
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        self.table.verticalHeader().setVisible(False)
        layout.addWidget(self.table)

        self.status_label = QLabel("")
        layout.addWidget(self.status_label)

    def update_data(self, yr: YearResult):
        self.setWindowTitle(self.tr("電源別サマリー — {}年").format(yr.year))
        self.title_label.setText(self.tr("{}年  ステータス: {}").format(yr.year, yr.status))

        all_c = sorted(set(list(yr.capacity_by_carrier) + list(yr.generation_by_carrier)))
        self.table.setRowCount(len(all_c))
        for row, c in enumerate(all_c):
            self.table.setItem(row, 0, QTableWidgetItem(c))
            self.table.setItem(row, 1, QTableWidgetItem(
                f"{yr.capacity_by_carrier.get(c, 0):,.1f}"))
            self.table.setItem(row, 2, QTableWidgetItem(
                f"{yr.generation_by_carrier.get(c, 0):,.0f}"))
            self.table.setItem(row, 3, QTableWidgetItem(
                f"{yr.capex_by_carrier.get(c, 0):,.0f}"))
            self.table.setItem(row, 4, QTableWidgetItem(
                f"{yr.opex_by_carrier.get(c, 0):,.0f}"))

        total_gen  = sum(yr.generation_by_carrier.values())
        total_cost = sum(yr.capex_by_carrier.values()) + sum(yr.opex_by_carrier.values())
        lcoe        = total_cost / total_gen if total_gen > 0 else 0.0
        co2_per_kwh = yr.co2_emissions * 1000 / total_gen if total_gen > 0 else 0.0
        self.status_label.setText(
            f"目的関数: {yr.objective:,.0f} Currency  |  CO₂: {yr.co2_emissions:,.0f} tCO₂  "
            f"|  LCOE: {lcoe:,.1f} Currency/MWh  |  CO₂原単位: {co2_per_kwh:,.1f} gCO₂/kWh"
        )


class ResultsPanel(QWidget):
    """Shown in the「結果」tab. Populated after optimization finishes."""

    def __init__(self, parent=None):
        super().__init__(parent)
        self._results: OptimizationResults | None = None
        self._all_scenarios: dict[str, OptimizationResults] = {}
        self._summary_win: _SummaryWindow | None = None
        self._current_yr: YearResult | None = None
        self._current_snapshot_step: int = 1
        self._default_results_dir = os.path.join(
            os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "results")
        self._setup_ui()

    # ── UI ────────────────────────────────────────────────────────────
    def _setup_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(4, 4, 4, 4)

        # Top bar
        top = QHBoxLayout()
        top.addWidget(QLabel(self.tr("シナリオ:")))
        self.scenario_combo = QComboBox()
        self.scenario_combo.setMinimumWidth(160)
        self.scenario_combo.currentIndexChanged.connect(self._on_scenario_changed)
        top.addWidget(self.scenario_combo)
        top.addSpacing(12)
        self.year_combo = QComboBox()
        self.year_combo.currentIndexChanged.connect(self._on_year_changed)
        top.addWidget(QLabel(self.tr("表示年:")))
        top.addWidget(self.year_combo)
        top.addStretch()
        btn_open_nc = QPushButton(self.tr("netCDFを開く..."))
        btn_open_nc.clicked.connect(self.open_netcdf_dialog)
        top.addWidget(btn_open_nc)
        btn_export = QPushButton(self.tr("Excelエクスポート"))
        btn_export.clicked.connect(self._export_excel)
        top.addWidget(btn_export)
        layout.addLayout(top)

        # No-data placeholder
        self.placeholder = QLabel(self.tr("最適化を実行すると結果がここに表示されます。"))
        self.placeholder.setAlignment(Qt.AlignmentFlag.AlignCenter)
        layout.addWidget(self.placeholder)

        # Tabs for different chart types
        self.tab_widget = QTabWidget()
        self.tab_widget.setVisible(False)

        # Tab 1: single-year summary
        self.single_tab = QWidget()
        self._build_single_tab()
        self.tab_widget.addTab(self.single_tab, self.tr("単年詳細"))

        # Tab 2: multi-year comparison
        self.multi_tab = QWidget()
        self._build_multi_tab()
        self.tab_widget.addTab(self.multi_tab, self.tr("多年比較"))

        layout.addWidget(self.tab_widget, stretch=1)

    def _build_single_tab(self):
        outer = QVBoxLayout(self.single_tab)
        outer.setContentsMargins(0, 0, 0, 0)
        outer.setSpacing(0)

        splitter = QSplitter(Qt.Orientation.Horizontal)
        splitter.setChildrenCollapsible(False)

        # ── Left pane: charts + table + status ───────────────────────
        left_widget = QWidget()
        left = QVBoxLayout(left_widget)
        left.setSpacing(4)
        left.setContentsMargins(4, 4, 2, 4)

        # Capacity chart
        self.cap_fig = Figure(figsize=(5, 3), dpi=90, tight_layout=True)
        self.cap_ax  = self.cap_fig.add_subplot(111)
        cap_canvas = FigureCanvas(self.cap_fig)
        cap_canvas.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding)
        left.addWidget(cap_canvas, stretch=1)

        # Generation chart
        self.gen_fig = Figure(figsize=(5, 3), dpi=90, tight_layout=True)
        self.gen_ax  = self.gen_fig.add_subplot(111)
        gen_canvas = FigureCanvas(self.gen_fig)
        gen_canvas.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding)
        left.addWidget(gen_canvas, stretch=1)

        # Generation mix pie
        self.cost_fig = Figure(figsize=(5, 3), dpi=90, tight_layout=True)
        self.cost_ax  = self.cost_fig.add_subplot(111)
        cost_canvas = FigureCanvas(self.cost_fig)
        cost_canvas.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding)
        left.addWidget(cost_canvas, stretch=1)

        # Summary table button
        btn_summary = QPushButton(self.tr("電源別サマリーを表示..."))
        btn_summary.clicked.connect(self._show_summary_window)
        left.addWidget(btn_summary)

        # Status bar
        self.status_label = QLabel("")
        left.addWidget(self.status_label)

        splitter.addWidget(left_widget)

        # ── Right pane: dispatch simulation ──────────────────────────
        right_widget = QWidget()
        right = QVBoxLayout(right_widget)
        right.setSpacing(4)
        right.setContentsMargins(2, 4, 4, 4)

        # Controls row
        ctrl = QHBoxLayout()
        ctrl.addWidget(QLabel(self.tr("開始 [h]:")))
        self.disp_start = QSpinBox()
        self.disp_start.setRange(0, 8759)
        self.disp_start.setValue(0)
        ctrl.addWidget(self.disp_start)

        self.disp_start_label = QLabel("")
        ctrl.addWidget(self.disp_start_label)

        ctrl.addWidget(QLabel(self.tr("  終了 [h]:")))
        self.disp_end = QSpinBox()
        self.disp_end.setRange(1, 8760)
        self.disp_end.setValue(168)
        ctrl.addWidget(self.disp_end)

        self.disp_end_label = QLabel("")
        ctrl.addWidget(self.disp_end_label)

        ctrl.addSpacing(12)
        for label, slot in [
            (self.tr("全期間"),   lambda: self._set_disp_range(0, self.disp_end.maximum())),
            (self.tr("1ヶ月"),   lambda: self._set_disp_range(self.disp_start.value(),
                                    self.disp_start.value() + max(1, 720 // self._current_snapshot_step))),
            (self.tr("1週間"),   lambda: self._set_disp_range(self.disp_start.value(),
                                    self.disp_start.value() + max(1, 168 // self._current_snapshot_step))),
            (self.tr("1日"),     lambda: self._set_disp_range(self.disp_start.value(),
                                    self.disp_start.value() + max(1, 24 // self._current_snapshot_step))),
            (self.tr("◄ 前へ"),  self._disp_prev),
            (self.tr("次へ ►"),  self._disp_next),
        ]:
            b = QPushButton(label)
            b.setFixedHeight(26)
            b.clicked.connect(slot)
            ctrl.addWidget(b)
        ctrl.addStretch()
        right.addLayout(ctrl)

        # Chart 1: 需給バランス
        self.disp_fig = Figure(figsize=(12, 4), dpi=90, tight_layout=True)
        self.disp_ax  = self.disp_fig.add_subplot(111)
        self.disp_canvas = FigureCanvas(self.disp_fig)
        self.disp_canvas.setSizePolicy(
            QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding)
        right.addWidget(self.disp_canvas, stretch=3)

        # Chart 2: 地域間潮流
        self.flow_fig = Figure(figsize=(12, 3), dpi=90, tight_layout=True)
        self.flow_ax  = self.flow_fig.add_subplot(111)
        self.flow_canvas = FigureCanvas(self.flow_fig)
        self.flow_canvas.setSizePolicy(
            QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding)
        right.addWidget(self.flow_canvas, stretch=2)

        # Chart 3: 再エネ出力抑制
        self.curt_fig = Figure(figsize=(12, 3), dpi=90, tight_layout=True)
        self.curt_ax  = self.curt_fig.add_subplot(111)
        self.curt_canvas = FigureCanvas(self.curt_fig)
        self.curt_canvas.setSizePolicy(
            QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding)
        right.addWidget(self.curt_canvas, stretch=2)

        # Connect spinboxes
        self.disp_start.valueChanged.connect(self._on_disp_range_changed)
        self.disp_end.valueChanged.connect(self._on_disp_range_changed)
        self._update_disp_date_labels()

        splitter.addWidget(right_widget)
        splitter.setSizes([400, 600])
        outer.addWidget(splitter)

    def _build_multi_tab(self):
        layout = QVBoxLayout(self.multi_tab)
        layout.setSpacing(4)

        # Capacity over years
        self.multi_cap_fig = Figure(figsize=(10, 3), dpi=90, tight_layout=True)
        self.multi_cap_ax  = self.multi_cap_fig.add_subplot(111)
        layout.addWidget(FigureCanvas(self.multi_cap_fig))

        # CO₂, cost, LCOE, CO₂ intensity over years
        row = QHBoxLayout()
        self.multi_co2_fig = Figure(figsize=(5, 3), dpi=90, tight_layout=True)
        self.multi_co2_ax  = self.multi_co2_fig.add_subplot(111)
        row.addWidget(FigureCanvas(self.multi_co2_fig))

        self.multi_cost_fig = Figure(figsize=(5, 3), dpi=90, tight_layout=True)
        self.multi_cost_ax  = self.multi_cost_fig.add_subplot(111)
        row.addWidget(FigureCanvas(self.multi_cost_fig))
        layout.addLayout(row)

        row2 = QHBoxLayout()
        self.multi_lcoe_fig = Figure(figsize=(5, 3), dpi=90, tight_layout=True)
        self.multi_lcoe_ax  = self.multi_lcoe_fig.add_subplot(111)
        row2.addWidget(FigureCanvas(self.multi_lcoe_fig))

        self.multi_co2int_fig = Figure(figsize=(5, 3), dpi=90, tight_layout=True)
        self.multi_co2int_ax  = self.multi_co2int_fig.add_subplot(111)
        row2.addWidget(FigureCanvas(self.multi_co2int_fig))
        layout.addLayout(row2)

    # ── Public API ────────────────────────────────────────────────────
    @staticmethod
    def _parse_nc_filename(basename: str) -> tuple[str, int]:
        """ファイル名からシナリオ名と年を解析する。
        例:
          result_2024.nc              -> ("ベースライン", 2024)
          result_Scenario1_2030.nc    -> ("Scenario1", 2030)
          result_Scenario1_2030_太陽光２倍.nc -> ("Scenario1", 2030)
          result_Scenario1_mp_2030-2040-2050_step24.nc -> ("Scenario1", 0)
        """
        stem = re.sub(r"\.nc$", "", basename, flags=re.IGNORECASE)
        if stem.startswith("result_"):
            rest = stem[len("result_"):]
        else:
            rest = stem
        # multi-period: <name>_mp_<years>_...
        m_mp = re.match(r"(.+?)_mp_(\d{4}(?:-\d{4})*)", rest)
        if m_mp:
            return m_mp.group(1), 0
        # ベースライン: 先頭が数字
        if rest and rest[0].isdigit():
            m = re.search(r"(\d{4})", rest)
            year = int(m.group(1)) if m else 0
            return "ベースライン", year
        # シナリオあり: <name>_<year>[_<description>]
        m = re.match(r"(.+?)_(\d{4})", rest)
        if m:
            return m.group(1), int(m.group(2))
        # フォールバック
        m2 = re.search(r"(\d{4})", rest)
        year = int(m2.group(1)) if m2 else 0
        return rest or "ベースライン", year

    def open_netcdf_dialog(self):
        """Open a multi-select file dialog and load netCDF result files."""
        os.makedirs(self._default_results_dir, exist_ok=True)
        paths, _ = QFileDialog.getOpenFileNames(
            self, self.tr("netCDF結果ファイルを開く"), self._default_results_dir,
            "NetCDF Files (*.nc);;All Files (*)",
        )
        if paths:
            self.load_from_netcdf(paths)

    def load_from_netcdf(self, filepaths: list):
        """Load one or more result_{year}.nc files and populate the results panel."""
        try:
            import pypsa
        except ImportError:
            QMessageBox.critical(self, self.tr("エラー"), "PyPSA がインストールされていません。")
            return

        # ファイルをシナリオ名でグルーピング
        grouped: dict[str, list[tuple[str, int]]] = {}
        for path in sorted(filepaths):
            basename = os.path.basename(path)
            scenario_name, year = self._parse_nc_filename(basename)
            grouped.setdefault(scenario_name, []).append((path, year))

        all_scenarios: dict[str, OptimizationResults] = {}
        all_errors: list = []

        for scenario_name, file_list in sorted(grouped.items()):
            results = OptimizationResults(scenario_name=scenario_name)
            for path, year in file_list:
                basename = os.path.basename(path)
                try:
                    tmp_fd, tmp_path = tempfile.mkstemp(suffix='.nc')
                    os.close(tmp_fd)
                    try:
                        shutil.copy2(path, tmp_path)
                        n = pypsa.Network(tmp_path)
                    finally:
                        try:
                            os.unlink(tmp_path)
                        except OSError:
                            pass

                    import pandas as pd
                    is_multi = (hasattr(n, "investment_periods")
                                and n.investment_periods is not None
                                and len(n.investment_periods) > 1)

                    if is_multi:
                        periods = sorted(n.investment_periods)
                        n_snaps_per_period = len(n.snapshots) // len(periods) if len(periods) else 1
                        step = max(1, 8760 // max(1, n_snaps_per_period))
                        obj = 0.0
                        try:
                            obj = float(n.objective)
                        except Exception:
                            pass
                        for period in periods:
                            yr = YearResult(year=int(period), snapshot_step=step)
                            yr.status = "ok"
                            yr.objective = obj / len(periods)
                            _extract_multi_period_year(n, yr, int(period))
                            results.year_results.append(yr)
                    else:
                        yr = YearResult(year=year)
                        yr.status = "ok"
                        yr.snapshot_step = max(1, 8760 // max(1, len(n.snapshots)))
                        try:
                            yr.objective = float(n.objective)
                        except Exception:
                            yr.objective = 0.0
                        extract_year_results(n, yr)
                        extract_year_timeseries(n, yr)
                        results.year_results.append(yr)
                except Exception:
                    yr = YearResult(year=year)
                    yr.status = "error"
                    all_errors.append(f"{basename}:\n{traceback.format_exc()}")
                    results.year_results.append(yr)
            results.year_results.sort(key=lambda r: r.year)
            all_scenarios[scenario_name] = results

        if all_errors:
            QMessageBox.warning(
                self, "読み込みエラー",
                "一部のファイルで読み込みに失敗しました:\n\n" + "\n\n".join(all_errors),
            )

        if all_scenarios:
            self._load_multi_scenarios(all_scenarios)

    def load_results(self, results: OptimizationResults):
        """最適化実行後に呼ばれる（単一シナリオ）。"""
        name = results.scenario_name or "デフォルト"
        self._load_multi_scenarios({name: results})

    def load_results_map(self, scenarios: dict[str, OptimizationResults]):
        """複数シナリオ結果をまとめてロードする。"""
        self._load_multi_scenarios(scenarios)

    def _load_multi_scenarios(self, scenarios: dict[str, OptimizationResults]):
        """複数シナリオの結果をまとめてロードする。"""
        self._all_scenarios = scenarios
        self.placeholder.setVisible(False)
        self.tab_widget.setVisible(True)

        self.scenario_combo.blockSignals(True)
        self.scenario_combo.clear()
        self.scenario_combo.addItems(list(scenarios.keys()))
        self.scenario_combo.blockSignals(False)

        # 先頭シナリオを表示
        if scenarios:
            first = list(scenarios.values())[0]
            self._apply_results(first)

    def _apply_results(self, results: OptimizationResults):
        """指定シナリオの結果を全タブに反映する。"""
        self._results = results
        years = [yr.year for yr in results.year_results]
        self.year_combo.blockSignals(True)
        self.year_combo.clear()
        self.year_combo.addItems([str(y) for y in years])
        self.year_combo.blockSignals(False)

        if years:
            self._sync_disp_controls(results.year_results[0])
            self._draw_single(results.year_results[0])
            self._draw_dispatch(results.year_results[0])
        self._draw_multi(results)

    # ── Drawing ───────────────────────────────────────────────────────
    def _on_scenario_changed(self, idx: int):
        name = self.scenario_combo.currentText()
        if name in self._all_scenarios:
            self._apply_results(self._all_scenarios[name])

    def _on_year_changed(self, idx: int):
        if not self._results or idx < 0:
            return
        if idx < len(self._results.year_results):
            yr = self._results.year_results[idx]
            self._sync_disp_controls(yr)
            self._draw_single(yr)
            self._draw_dispatch(yr)

    def _draw_single(self, yr: YearResult):
        colors = [CARRIER_COLORS.get(c, "#808080") for c in yr.capacity_by_carrier]

        # Capacity bar
        self.cap_ax.clear()
        if yr.capacity_by_carrier:
            carriers = list(yr.capacity_by_carrier.keys())
            vals     = [yr.capacity_by_carrier[c] for c in carriers]
            self.cap_ax.bar(carriers, vals, color=colors)
            self.cap_ax.set_title(self.tr("最適設備容量 ({year}年) [MW]").format(year=yr.year))
            self.cap_ax.tick_params(axis="x", rotation=45)
        self.cap_fig.canvas.draw()

        # Generation bar
        self.gen_ax.clear()
        if yr.generation_by_carrier:
            carriers = list(yr.generation_by_carrier.keys())
            vals     = [yr.generation_by_carrier[c] / 1e6 for c in carriers]
            colors2  = [CARRIER_COLORS.get(c, "#808080") for c in carriers]
            self.gen_ax.bar(carriers, vals, color=colors2)
            self.gen_ax.set_title(self.tr("発電量 ({year}年) [TWh]").format(year=yr.year))
            self.gen_ax.tick_params(axis="x", rotation=45)
        self.gen_fig.canvas.draw()

        # Generation mix pie
        self.cost_ax.clear()
        pos_g = [(c, v) for c, v in yr.generation_by_carrier.items() if v > 0]
        if pos_g:
            labels, vals = zip(*pos_g)
            pie_colors = [CARRIER_COLORS.get(c, "#808080") for c in labels]
            self.cost_ax.pie(vals, labels=labels, colors=pie_colors,
                             autopct="%1.1f%%", startangle=90)
            self.cost_ax.set_title(self.tr("発電量内訳 ({year}年)").format(year=yr.year))
        self.cost_fig.canvas.draw()

        # Table — ウィンドウが開いていれば自動更新
        self._current_yr = yr
        if self._summary_win is not None and self._summary_win.isVisible():
            self._summary_win.update_data(yr)

        total_gen  = sum(yr.generation_by_carrier.values())
        total_cost = sum(yr.capex_by_carrier.values()) + sum(yr.opex_by_carrier.values())
        lcoe       = total_cost / total_gen if total_gen > 0 else 0.0
        co2_per_kwh = yr.co2_emissions * 1000 / total_gen if total_gen > 0 else 0.0

        # 再エネ出力抑制率 = 抑制量 / (抑制量 + 実発電量) × 100
        total_curt = (
            float(yr.curtailment_df.values.sum())
            if yr.curtailment_df is not None else 0.0
        )
        total_re_gen = sum(
            yr.generation_by_carrier.get(c, 0.0) for c in CF_CARRIERS
        )
        re_potential = total_curt + total_re_gen
        curt_rate = total_curt / re_potential * 100.0 if re_potential > 0 else 0.0

        self.status_label.setText(
            f"ステータス: {yr.status}  |  目的関数: {yr.objective:,.0f} Currency  "
            f"|  CO₂排出量: {yr.co2_emissions:,.0f} tCO₂  "
            f"|  LCOE: {lcoe:,.1f} Currency/MWh  |  CO₂原単位: {co2_per_kwh:,.1f} gCO₂/kWh  "
            f"|  再エネ出力抑制率: {curt_rate:.1f}%"
        )

    def _show_summary_window(self):
        """電源別サマリーウィンドウを開く（なければ作成、既存なら前面に出す）。"""
        if self._summary_win is None or not self._summary_win.isVisible():
            self._summary_win = _SummaryWindow()
        yr = getattr(self, "_current_yr", None)
        if yr is not None:
            self._summary_win.update_data(yr)
        self._summary_win.show()
        self._summary_win.raise_()
        self._summary_win.activateWindow()

    def _draw_multi(self, results: OptimizationResults):
        if not results.year_results:
            return
        years = [yr.year for yr in results.year_results]

        # Multi-year capacity (stacked bar)
        all_carriers = sorted({c for yr in results.year_results
                                for c in yr.capacity_by_carrier})
        self.multi_cap_ax.clear()
        bottoms = np.zeros(len(years))
        for carrier in all_carriers:
            vals  = np.array([yr.capacity_by_carrier.get(carrier, 0)
                              for yr in results.year_results])
            color = CARRIER_COLORS.get(carrier, "#808080")
            self.multi_cap_ax.bar(years, vals, bottom=bottoms, label=carrier, color=color)
            bottoms += vals
        self.multi_cap_ax.set_title(self.tr("設備容量推移 [MW]"))
        self.multi_cap_ax.set_xlabel(self.tr("計画年"))
        handles, labels = self.multi_cap_ax.get_legend_handles_labels()
        if handles and labels:
            self.multi_cap_ax.legend(fontsize=7, ncol=3, loc="upper left")
        self.multi_cap_fig.canvas.draw()

        # CO₂
        co2_vals = [yr.co2_emissions for yr in results.year_results]
        self.multi_co2_ax.clear()
        self.multi_co2_ax.plot(years, co2_vals, marker="o", color="#e74c3c")
        self.multi_co2_ax.set_title(self.tr("CO₂排出量推移 [tCO₂]"))
        self.multi_co2_ax.set_xlabel(self.tr("計画年"))
        self.multi_co2_fig.canvas.draw()

        # Total cost
        cost_vals = [yr.objective for yr in results.year_results]
        self.multi_cost_ax.clear()
        self.multi_cost_ax.plot(years, cost_vals, marker="o", color="#2980b9")
        self.multi_cost_ax.set_title(self.tr("総コスト推移 [Currency]"))
        self.multi_cost_ax.set_xlabel(self.tr("計画年"))
        self.multi_cost_fig.canvas.draw()

        # LCOE and CO₂ intensity
        lcoe_vals    = []
        co2int_vals  = []
        for yr in results.year_results:
            total_gen  = sum(yr.generation_by_carrier.values())
            total_cost = sum(yr.capex_by_carrier.values()) + sum(yr.opex_by_carrier.values())
            lcoe_vals.append(total_cost / total_gen if total_gen > 0 else 0.0)
            co2int_vals.append(yr.co2_emissions * 1000 / total_gen if total_gen > 0 else 0.0)

        self.multi_lcoe_ax.clear()
        self.multi_lcoe_ax.plot(years, lcoe_vals, marker="o", color="#27ae60")
        self.multi_lcoe_ax.set_title(self.tr("LCOE推移 [Currency/MWh]"))
        self.multi_lcoe_ax.set_xlabel(self.tr("計画年"))
        self.multi_lcoe_fig.canvas.draw()

        self.multi_co2int_ax.clear()
        self.multi_co2int_ax.plot(years, co2int_vals, marker="o", color="#e67e22")
        self.multi_co2int_ax.set_title(self.tr("CO₂原単位推移 [gCO₂/kWh]"))
        self.multi_co2int_ax.set_xlabel(self.tr("計画年"))
        self.multi_co2int_fig.canvas.draw()

    # ── Dispatch simulation ───────────────────────────────────────────
    _BASE_DATE = datetime.datetime(2019, 1, 1)

    @staticmethod
    def _series_len(yr: YearResult) -> int:
        lengths: list[int] = []
        for obj in (
            yr.dispatch_df,
            yr.link_gen_df,
            yr.link_charge_df,
            yr.link_flow_df,
            yr.curtailment_df,
            yr.demand_ts,
        ):
            if obj is None:
                continue
            try:
                lengths.append(len(obj))
            except TypeError:
                pass
        return max(lengths) if lengths else 0

    @staticmethod
    def _slice_1d(values, start: int, end: int) -> np.ndarray:
        arr = np.asarray(values)
        if arr.ndim == 0:
            return np.array([], dtype=float)
        return arr[start:end]

    def _sync_disp_controls(self, yr: YearResult):
        self._current_snapshot_step = max(1, getattr(yr, 'snapshot_step', 1))
        max_points = max(1, self._series_len(yr))
        start = min(self.disp_start.value(), max_points - 1)
        end = min(self.disp_end.value(), max_points)
        if end <= start:
            end = min(max_points, start + 1)

        self.disp_start.blockSignals(True)
        self.disp_end.blockSignals(True)
        self.disp_start.setRange(0, max_points - 1)
        self.disp_end.setRange(1, max_points)
        self.disp_start.setValue(start)
        self.disp_end.setValue(end)
        self.disp_start.blockSignals(False)
        self.disp_end.blockSignals(False)
        self._update_disp_date_labels()

    def _hour_to_date(self, h: int) -> str:
        dt = self._BASE_DATE + datetime.timedelta(hours=int(h) * self._current_snapshot_step)
        return dt.strftime("%m/%d %H:00")

    def _update_disp_date_labels(self):
        self.disp_start_label.setText(f"({self._hour_to_date(self.disp_start.value())})")
        max_end_hour = max(0, self.disp_end.maximum() - 1)
        self.disp_end_label.setText(
            f"({self._hour_to_date(min(self.disp_end.value(), max_end_hour))})")

    def _on_disp_range_changed(self):
        self._update_disp_date_labels()
        if not self._results:
            return
        idx = self.year_combo.currentIndex()
        if 0 <= idx < len(self._results.year_results):
            self._draw_dispatch(self._results.year_results[idx])

    def _set_disp_range(self, start: int, end: int):
        max_end = self.disp_end.maximum()
        w = max(1, end - start)                  # ウィンドウ幅を保持
        w = min(w, max_end)
        start = max(0, min(start, max_end - w))  # 端に当たったら start を手前にスライド
        end = start + w
        self.disp_start.blockSignals(True)
        self.disp_end.blockSignals(True)
        self.disp_start.setValue(start)
        self.disp_end.setValue(end)
        self.disp_start.blockSignals(False)
        self.disp_end.blockSignals(False)
        self._on_disp_range_changed()

    def _disp_prev(self):
        w = self.disp_end.value() - self.disp_start.value()
        self._set_disp_range(self.disp_start.value() - w, self.disp_start.value())

    def _disp_next(self):
        w = self.disp_end.value() - self.disp_start.value()
        self._set_disp_range(self.disp_end.value(), self.disp_end.value() + w)

    def _draw_dispatch(self, yr: YearResult):
        max_points = self._series_len(yr)
        if max_points <= 0:
            return

        start = min(self.disp_start.value(), max_points - 1)
        end = min(self.disp_end.value(), max_points)
        if start >= end:
            return

        h = np.arange(start, end)
        range_str = f"{self._hour_to_date(start)} 〜 {self._hour_to_date(end)}"

        _step = self._current_snapshot_step
        def _fmt_hour(x, _):
            hi = int(x)
            actual_hour = hi * _step
            if 0 <= actual_hour <= 8760:
                dt = self._BASE_DATE + datetime.timedelta(hours=actual_hour)
                return dt.strftime("%m/%d\n%H:00")
            return ""

        _link_display = {"Water": "揚水", "heat": "熱", "hydrogen": "水素", "gas": "ガス", "DC": "DC"}
        _STACK_ORDER = [
            "Nuclear", "Hydro", "Biomass", "Wind",
            "Coal", "Gas", "Oil", "Solar", "揚水放電", "Other",
        ]
        def _order_key(label):
            try:
                return _STACK_ORDER.index(label)
            except ValueError:
                return len(_STACK_ORDER)

        # ── Chart 1: 需給バランス ─────────────────────────────────────
        self.disp_fig.clear()
        ax = self.disp_fig.add_subplot(111)
        self.disp_ax = ax

        stack_labels: list = []
        stack_data:   list = []
        stack_colors: list = []
        stor_charge_bands: list = []

        if yr.dispatch_df is not None:
            for carrier in yr.dispatch_df.columns:
                arr = np.maximum(self._slice_1d(yr.dispatch_df[carrier].values, start, end), 0.0)
                if arr.sum() > 0:
                    stack_labels.append(carrier)
                    stack_data.append(arr)
                    stack_colors.append(CARRIER_COLORS.get(carrier, "#808080"))

        if yr.link_gen_df is not None:
            for carrier in yr.link_gen_df.columns:
                arr = np.maximum(self._slice_1d(yr.link_gen_df[carrier].values, start, end), 0.0)
                if arr.sum() > 0:
                    name = _link_display.get(carrier, carrier)
                    stack_labels.append(f"{name}放電")
                    stack_data.append(arr)
                    stack_colors.append(CARRIER_COLORS.get(carrier, "#1E90FF"))

        if yr.link_charge_df is not None:
            for carrier in yr.link_charge_df.columns:
                arr = np.minimum(self._slice_1d(yr.link_charge_df[carrier].values, start, end), 0.0)
                if arr.sum() < 0:
                    name = _link_display.get(carrier, carrier)
                    stor_charge_bands.append(
                        (f"{name}充電", CARRIER_COLORS.get(carrier, "#1E90FF"), arr))

        if stack_data:
            combined = sorted(
                zip(stack_labels, stack_data, stack_colors),
                key=lambda t: _order_key(t[0]),
            )
            sorted_labels, sorted_data, sorted_colors = zip(*combined)
            m = min([len(h)] + [len(a) for a in sorted_data])
            if m > 0:
                x_stack = h[:m]
                y_stack = [a[:m] for a in sorted_data]
                ax.stackplot(x_stack, *y_stack, labels=sorted_labels,
                             colors=sorted_colors, alpha=0.85)

        # ── 需要（ゼロ軸より下）と蓄電・揚水の充電（さらに下に積み上げ）──
        neg_bottom = np.zeros(len(h))

        if yr.demand_ts is not None:
            demand = self._slice_1d(yr.demand_ts.values, start, end)
            m = min(len(h), len(demand))
            if m > 0:
                demand_neg = -np.maximum(demand[:m], 0.0)
                ax.fill_between(h[:m], neg_bottom[:m], neg_bottom[:m] + demand_neg,
                                alpha=0.85, color="firebrick", label="需要")
                neg_bottom[:m] += demand_neg

        for label, color, arr in stor_charge_bands:
            m = min(len(h), len(arr))
            if m <= 0:
                continue
            ax.fill_between(h[:m], neg_bottom[:m], neg_bottom[:m] + arr[:m],
                            alpha=0.85, color=color, label=label)
            neg_bottom[:m] += arr[:m]

        ax.axhline(0, color="black", linewidth=0.5, zorder=5)
        ax.xaxis.set_major_formatter(FuncFormatter(_fmt_hour))
        ax.set_xlim(start, end)
        ax.set_ylabel("電力 [MW]")
        ax.set_title(f"需給バランス ({yr.year}年) — {range_str}")

        has_balance = any(x is not None for x in
                          (yr.dispatch_df, yr.link_gen_df, yr.link_charge_df, yr.demand_ts))
        if has_balance and (stack_labels or stor_charge_bands or yr.demand_ts is not None):
            ax.legend(fontsize=7, ncol=4, loc="upper right", framealpha=0.8)
        elif not has_balance:
            ax.text(0.5, 0.5, "時系列データがありません\n（最適化後に自動取得されます）",
                    transform=ax.transAxes, ha="center", va="center", fontsize=11)
        self.disp_canvas.draw()

        # ── Chart 2: 地域間潮流 ───────────────────────────────────────
        self.flow_fig.clear()
        ax2 = self.flow_fig.add_subplot(111)
        self.flow_ax = ax2

        has_flow = (yr.link_flow_df is not None and not yr.link_flow_df.empty)
        if has_flow:
            flow_colors = plt.cm.tab10(np.linspace(0, 0.9, len(yr.link_flow_df.columns)))
            for i, lk_name in enumerate(yr.link_flow_df.columns):
                arr = self._slice_1d(yr.link_flow_df[lk_name].values, start, end)
                m = min(len(h), len(arr))
                if m > 0:
                    ax2.plot(h[:m], arr[:m], linewidth=1.2, color=flow_colors[i],
                             label=lk_name, zorder=8)
            ax2.axhline(0, color="gray", linewidth=0.4, linestyle=":", zorder=3)
            ax2.legend(fontsize=7, ncol=2, loc="upper right", framealpha=0.7)
        else:
            ax2.text(0.5, 0.5, "潮流データなし", transform=ax2.transAxes,
                     ha="center", va="center", fontsize=10, color="gray")

        ax2.xaxis.set_major_formatter(FuncFormatter(_fmt_hour))
        ax2.set_xlim(start, end)
        ax2.set_ylabel("潮流 [MW]")
        ax2.set_title(f"地域間潮流 ({yr.year}年) — {range_str}")
        self.flow_canvas.draw()

        # ── Chart 3: 再エネ出力抑制 ───────────────────────────────────
        self.curt_fig.clear()
        ax3 = self.curt_fig.add_subplot(111)
        self.curt_ax = ax3

        curt_bottom = np.zeros(len(h))
        drawn_curt = False
        if yr.curtailment_df is not None:
            for carrier in yr.curtailment_df.columns:
                arr = self._slice_1d(yr.curtailment_df[carrier].values, start, end)
                if arr.sum() > 0:
                    color = CARRIER_COLORS.get(carrier, "#808080")
                    m = min(len(h), len(arr))
                    if m <= 0:
                        continue
                    ax3.fill_between(h[:m], curt_bottom[:m], curt_bottom[:m] + arr[:m],
                                     alpha=0.7, color=color, hatch="//",
                                     label=f"{carrier}出力抑制")
                    curt_bottom[:m] = curt_bottom[:m] + arr[:m]
                    drawn_curt = True

        if drawn_curt:
            ax3.legend(fontsize=7, loc="upper right", framealpha=0.7)
        else:
            ax3.text(0.5, 0.5, "出力抑制なし", transform=ax3.transAxes,
                     ha="center", va="center", fontsize=10, color="gray")

        ax3.xaxis.set_major_formatter(FuncFormatter(_fmt_hour))
        ax3.set_xlim(start, end)
        ax3.set_ylabel("出力抑制 [MW]")
        ax3.set_title(f"再エネ出力抑制 ({yr.year}年) — {range_str}")
        self.curt_canvas.draw()

    # ── Excel export ──────────────────────────────────────────────────
    def _export_excel(self):
        if not self._results or not self._results.year_results:
            QMessageBox.information(self, self.tr("情報"), self.tr("結果がありません。"))
            return
        path, _ = QFileDialog.getSaveFileName(
            self, self.tr("結果をExcelエクスポート"), "results.xlsx",
            "Excel Files (*.xlsx);;All Files (*)")
        if not path:
            return
        try:
            _export_results_excel(self._results, path)
            QMessageBox.information(self, self.tr("完了"), self.tr("エクスポート完了:\n") + path)
        except Exception as e:
            QMessageBox.critical(self, self.tr("エラー"), self.tr("エクスポート失敗:\n") + str(e))


def _safe(v):
    """Convert a value to an Excel-safe Python scalar (no nan/inf/numpy types)."""
    import math
    if v is None:
        return None
    try:
        f = float(v)
        if math.isnan(f) or math.isinf(f):
            return None
        return f
    except (TypeError, ValueError):
        return str(v)


def _export_results_excel(results: OptimizationResults, path: str):
    import pandas as pd
    from openpyxl import Workbook

    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "summary"

    # Summary sheet: one row per year
    headers = ["year", "status", "objective_Currency", "co2_tCO2"]
    ws_summary.append(headers)
    for yr in results.year_results:
        ws_summary.append([yr.year, yr.status, _safe(yr.objective), _safe(yr.co2_emissions)])

    # Per-year detail sheets
    for yr in results.year_results:
        ws = wb.create_sheet(title=f"year_{yr.year}")
        ws.append(["carrier", "capacity_MW", "generation_MWh", "capex_Currency", "opex_Currency"])
        all_c = sorted(set(list(yr.capacity_by_carrier) + list(yr.generation_by_carrier)))
        for c in all_c:
            ws.append([
                c,
                _safe(yr.capacity_by_carrier.get(c, 0)),
                _safe(yr.generation_by_carrier.get(c, 0)),
                _safe(yr.capex_by_carrier.get(c, 0)),
                _safe(yr.opex_by_carrier.get(c, 0)),
            ])

    # Multi-year capacity sheet
    ws_cap = wb.create_sheet("capacity_trend")
    all_c = sorted({c for yr in results.year_results for c in yr.capacity_by_carrier})
    ws_cap.append(["year"] + list(all_c))
    for yr in results.year_results:
        ws_cap.append([yr.year] + [_safe(yr.capacity_by_carrier.get(c, 0)) for c in all_c])

    # CO₂ & cost trend sheet
    ws_trend = wb.create_sheet("co2_cost_trend")
    ws_trend.append(["year", "co2_tCO2", "total_cost_Currency"])
    for yr in results.year_results:
        ws_trend.append([yr.year, _safe(yr.co2_emissions), _safe(yr.objective)])

    # Dispatch timeseries sheet (generator dispatch by carrier)
    for yr in results.year_results:
        if yr.dispatch_df is not None and not yr.dispatch_df.empty:
            ws_d = wb.create_sheet(title=f"dispatch_{yr.year}")
            df = yr.dispatch_df
            ws_d.append(["timestamp"] + list(df.columns))
            for ts, row in df.iterrows():
                ws_d.append([str(ts)] + [_safe(v) for v in row])

    # Log sheet
    ws_log = wb.create_sheet("log")
    for line in (results.log or "").splitlines():
        ws_log.append([line])

    wb.save(path)
